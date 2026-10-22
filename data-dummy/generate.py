"""Generator data dummy audit penjualan tiket group order dan refund HST.

Menghasilkan semua sumber data audit untuk periode 2025-01 s.d. 2025-03 ke folder
output/, lengkap dengan anomali terkontrol dan kuncinya (anomalies_truth.csv).
Seed tetap, jadi hasil selalu sama setiap kali dijalankan.

Pemakaian:
    python generate.py [--out output] [--seed 20250101]
"""
from __future__ import annotations

import argparse
import calendar
import csv
import hashlib
import random
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pandas as pd
from faker import Faker
from openpyxl import Workbook

SEED = 20250101
PERIODS = ["2025-01", "2025-02", "2025-03"]
MONTH_ID = {1: "JAN", 2: "FEB", 3: "MAR"}
MONTH_NAME = {1: "Januari", 2: "Februari", 3: "Maret"}

STATIONS = {
    "STA": {"name": "Stasiun A", "business_area": "BA01", "office_no": "LK-A01"},
    "STB": {"name": "Stasiun B", "business_area": "BA02", "office_no": "LK-B01"},
    "STC": {"name": "Stasiun C", "business_area": "BA03", "office_no": "LK-C01"},
}
# Tarif dasar per kelas (premium economy, business, first) dan durasi perjalanan (menit).
ROUTES = {
    ("STA", "STB"): ((150_000, 250_000, 400_000), 25),
    ("STB", "STC"): ((100_000, 175_000, 300_000), 15),
    ("STA", "STC"): ((250_000, 450_000, 600_000), 40),
}
SEAT_CLASSES = ["PREMIUM_ECONOMY", "BUSINESS", "FIRST"]
TAX_RATE = 0.11
REFUND_RATE = 0.75          # kebijakan: refund maksimal 75% harga tiket
CASH_REFUND_LIMIT = 300_000  # kebijakan: refund tunai hanya WNI dan maksimal Rp300.000
GATEWAY = {"DEBIT_CARD": "EDC-01", "CREDIT_CARD": "PG-BETA", "QRIS": "PG-ALFA",
           "E_WALLET": "PG-BETA", "VA": "PG-ALFA"}
BANKS = ["Bank Nusantara", "Bank Merdeka", "Bank Sejahtera", "Bank Samudra"]
FOREIGN = [("en_US", "USA"), ("en_GB", "GBR"), ("en_AU", "AUS"), ("de_DE", "DEU"), ("fr_FR", "FRA")]
COMPANY_ACCOUNT = "1020304050"


def half_up(x: float) -> int:
    return int(x + 0.5)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def rupiah_id(n: int) -> str:
    """Format angka gaya rekening koran: 1.234.000,00"""
    return f"{n:,.2f}".replace(",", "#").replace(".", ",").replace("#", ".")


class Generator:
    def __init__(self, out: Path, seed: int):
        self.out = out
        self.rng = random.Random(seed)
        Faker.seed(seed)
        self.fake = Faker("id_ID")
        self.foreign_fakers = {loc: Faker(loc) for loc, _ in FOREIGN}
        self.used_names: set[str] = set()
        self.ticket_seq = 0
        self.order_seq = 0
        self.truth: list[dict] = []
        self.decoys: list[dict] = []
        self.injected_tickets: set[str] = set()
        self.injected_refund_tickets: set[str] = set()

    # ------------------------------------------------------------------ helpers
    def unique_name(self) -> str:
        while True:
            name = f"{self.fake.first_name()} {self.fake.last_name()}"
            if name not in self.used_names:
                self.used_names.add(name)
                return name

    def nik(self) -> str:
        r = self.rng
        dob = date(r.randint(1960, 2005), r.randint(1, 12), r.randint(1, 28))
        return f"{r.randint(11, 94):02d}{r.randint(1, 79):02d}{r.randint(1, 40):02d}{dob:%d%m%y}{r.randint(1, 9999):04d}"

    def passenger(self) -> dict:
        if self.rng.random() < 0.92:
            return {"name": self.unique_name(), "id_no": self.nik(), "nationality": "IDN"}
        loc, nat = self.rng.choice(FOREIGN)
        f = self.foreign_fakers[loc]
        return {"name": f"{f.first_name()} {f.last_name()}",
                "id_no": f"{chr(65 + self.rng.randint(0, 25))}{self.rng.randint(10**7, 10**8 - 1)}",
                "nationality": nat}

    def add_truth(self, rule: str, entity_type: str, entity_id: str, dept: str, period: str, desc: str):
        self.truth.append({"anomaly_id": f"AN{len(self.truth) + 1:03d}", "rule_code": rule,
                           "entity_type": entity_type, "entity_id": entity_id,
                           "source_dept": dept, "period": period, "description": desc})

    def add_decoy(self, target_rule: str, entity_type: str, entity_id: str, desc: str):
        """Kasus pengecoh: mirip anomali tetapi sah, sehingga tidak boleh ditandai aturan mana pun."""
        self.decoys.append({"decoy_id": f"DC{len(self.decoys) + 1:03d}", "target_rule": target_rule,
                            "entity_type": entity_type, "entity_id": entity_id, "description": desc})

    @staticmethod
    def days(period: str) -> list[date]:
        y, m = map(int, period.split("-"))
        return [date(y, m, d) for d in range(1, calendar.monthrange(y, m)[1] + 1)]

    @staticmethod
    def period_of(d: date) -> str:
        return f"{d:%Y-%m}"

    @staticmethod
    def shift_of(t: time) -> int:
        return 1 if t < time(13, 30) else 2

    # -------------------------------------------------------------------- HCM
    def build_hcm(self):
        emps = []

        def add(dept, station, position):
            emp = {"employee_id": f"E{len(emps) + 1:04d}", "employee_name": self.unique_name(),
                   "department": dept, "station_code": station, "position": position,
                   "join_date": date(self.rng.randint(2019, 2024), self.rng.randint(1, 12), 1).isoformat(),
                   "employment_status": "AKTIF"}
            emps.append(emp)
            return emp

        self.station_staff = {}
        for st in STATIONS:
            self.station_staff[st] = {
                "kepala": add("STASIUN", st, "KEPALA_STASIUN"),
                "supervisor": [add("STASIUN", st, "SUPERVISOR") for _ in range(3)],
                "petugas": [add("STASIUN", st, "PETUGAS_LOKET") for _ in range(8)],
            }
        self.finance_mgr = add("FINANCE", None, "MANAGER_FINANCE")
        self.finance_staff = [add("FINANCE", None, "STAF_FINANCE") for _ in range(3)]
        self.it_mgr = add("IT", None, "MANAGER_IT")
        self.it_staff = [add("IT", None, "STAF_IT") for _ in range(2)]
        add("AUDIT_INTERNAL", None, "AUDITOR")
        add("AUDIT_INTERNAL", None, "AUDITOR")
        self.employees = emps

        roster = []
        self.roster = {}
        all_days = [d for p in PERIODS for d in self.days(p)]
        for di, d in enumerate(all_days):
            for st, staff in self.station_staff.items():
                for shift in (1, 2):
                    pet = staff["petugas"]
                    w1 = pet[(di * 4 + shift * 2) % 8]
                    w2 = pet[(di * 4 + shift * 2 + 1) % 8]
                    sup = staff["supervisor"][(di + shift) % 3]
                    self.roster[(d, st, shift)] = {1: w1, 2: w2, "sup": sup}
                    for role, emp, win in (("PETUGAS_LOKET", w1, 1), ("PETUGAS_LOKET", w2, 2), ("SUPERVISOR", sup, None)):
                        roster.append({"roster_date": d.isoformat(), "station_code": st, "shift_no": shift,
                                       "shift_start": "05:30" if shift == 1 else "13:30",
                                       "shift_end": "13:30" if shift == 1 else "21:30",
                                       "employee_id": emp["employee_id"], "role": role,
                                       "window_no": win})
        self.roster_rows = roster

    # ---------------------------------------------------------------- tickets
    def new_order_no(self, d: date) -> str:
        self.order_seq += 1
        return f"ORD{d:%Y%m%d}{self.order_seq:05d}"

    def make_ticket(self, *, d: date, t: time, origin: str, dest: str, channel: str, pay: str,
                    order_no: str, pax: dict, operator: str, shift: int | None, window: int | None,
                    station: str, discount_type: str = "NONE", discount_rate: int = 0,
                    seat_class: str | None = None, dep: tuple[date, time] | None = None,
                    bak_id: int | None = None, partner: str | None = None, trade: tuple | None = None) -> dict:
        self.ticket_seq += 1
        key = (origin, dest) if (origin, dest) in ROUTES else (dest, origin)
        fares, minutes = ROUTES[key]
        if seat_class is None:
            seat_class = self.rng.choices(SEAT_CLASSES, weights=[80, 15, 5])[0]
        original = fares[SEAT_CLASSES.index(seat_class)]
        before = half_up(original * (1 - discount_rate / 100))
        after = half_up(before * (1 + TAX_RATE))
        if dep is None:
            dep_d = d + timedelta(days=self.rng.choice([0, 0, 0, 0, 1, 2, 3, 7]))
            hour = self.rng.randint(6, 21) if dep_d > d else self.rng.randint(min(t.hour + 1, 21), 21)
            dep = (dep_d, time(hour, self.rng.choice([0, 30])))
        dep_dt = datetime.combine(*dep)
        arr = dep_dt + timedelta(minutes=minutes)
        if trade is None and pay in GATEWAY:
            trade = (f"T{self.rng.randint(10**11, 10**12 - 1)}", f"P{self.rng.randint(10**13, 10**14 - 1)}")
        return {
            "seq_no": 0, "periode": f"{MONTH_ID[d.month]}-{d.year}",
            "passenger_name": pax["name"], "nik_passport_no": pax["id_no"], "bak": bak_id,
            "nationality": pax["nationality"], "order_no": order_no, "ticket_no": f"TK{self.ticket_seq:010d}",
            "ticketing_station": station, "bussiness_area": STATIONS[origin]["business_area"],
            "office_no": STATIONS[station]["office_no"] if channel in ("COUNTER", "B2B") else None,
            "window_no": window, "shift_no": shift, "operator_name": operator,
            "ticketing_time": f"{d.isoformat()} {t:%H:%M:%S}", "departure_date": dep[0].isoformat(),
            "train_no": f"HS{(dep[1].hour - 5) * 2 + (1 if origin < dest else 2):03d}",
            "origin": STATIONS[origin]["name"], "cars_number": self.rng.randint(1, 8),
            "seat_number": f"{self.rng.randint(1, 20):02d}{self.rng.choice('ABCDF')}",
            "origin_code": origin, "purchase_date": d.isoformat(), "purchase_time": f"{t:%H:%M:%S}",
            "departure_time": f"{dep[1]:%H:%M}", "destination": STATIONS[dest]["name"],
            "destination_code": dest, "arrival_date": arr.date().isoformat(), "arrival_time": f"{arr:%H:%M}",
            "seat_class": seat_class, "ticket_type": "CHILD" if self.rng.random() < 0.05 else "ADULT",
            "original_ticket_price": original, "discount_type": discount_type, "discount_rate": discount_rate,
            "before_tax_price": before, "tax_rate": TAX_RATE, "after_tax_price": after,
            "ticketing_channel": channel, "payment_method": pay,
            "trade_no": trade[0] if trade else None, "plat_trade_no": trade[1] if trade else None,
            "payment_gateway": GATEWAY.get(pay), "b2b_partner": partner,
        }

    def rand_time(self, shift: int) -> time:
        start = datetime(2000, 1, 1, 5, 30) if shift == 1 else datetime(2000, 1, 1, 13, 30)
        return (start + timedelta(seconds=self.rng.randint(0, 8 * 3600 - 1))).time()

    def build_tickets(self):
        r = self.rng
        tickets = []
        for p in PERIODS:
            for d in self.days(p):
                for st in STATIONS:
                    others = [s for s in STATIONS if s != st]
                    # Loket stasiun
                    for _ in range(r.randint(5, 10)):
                        shift = r.choice((1, 2))
                        window = r.choice((1, 2))
                        t = self.rand_time(shift)
                        op = self.roster[(d, st, shift)][window]["employee_name"]
                        order = self.new_order_no(d)
                        pay = r.choices(["CASH", "DEBIT_CARD", "QRIS"], weights=[45, 30, 25])[0]
                        dest = r.choice(others)
                        dep = None
                        for i in range(r.randint(1, 4)):
                            tk = self.make_ticket(d=d, t=t, origin=st, dest=dest, channel="COUNTER", pay=pay,
                                                  order_no=order, pax=self.passenger(), operator=op, shift=shift,
                                                  window=window, station=st, dep=dep)
                            dep = (date.fromisoformat(tk["departure_date"]), time.fromisoformat(tk["departure_time"]))
                            tickets.append(tk)
                    # Mesin tiket (TVM)
                    for _ in range(r.randint(2, 5)):
                        t = self.rand_time(r.choice((1, 2)))
                        order = self.new_order_no(d)
                        pay = r.choice(["QRIS", "DEBIT_CARD"])
                        tickets.append(self.make_ticket(d=d, t=t, origin=st, dest=r.choice(others), channel="TVM",
                                                        pay=pay, order_no=order, pax=self.passenger(),
                                                        operator="SYSTEM", shift=self.shift_of(t), window=None,
                                                        station=st))
                    # Aplikasi dan web
                    for _ in range(r.randint(6, 12)):
                        t = (datetime(2000, 1, 1) + timedelta(seconds=r.randint(0, 86399))).time()
                        order = self.new_order_no(d)
                        channel = r.choice(["APP", "WEB"])
                        pay = r.choices(["E_WALLET", "VA", "CREDIT_CARD", "QRIS"], weights=[40, 25, 15, 20])[0]
                        promo = r.random() < 0.2
                        dest = r.choice(others)
                        dep = None
                        for i in range(r.randint(1, 3)):
                            tk = self.make_ticket(d=d, t=t, origin=st, dest=dest, channel=channel, pay=pay,
                                                  order_no=order, pax=self.passenger(), operator="SYSTEM",
                                                  shift=None, window=None, station="ONLINE",
                                                  discount_type="PROMO" if promo else "NONE",
                                                  discount_rate=10 if promo else 0, dep=dep)
                            dep = (date.fromisoformat(tk["departure_date"]), time.fromisoformat(tk["departure_time"]))
                            tickets.append(tk)
        self.tickets = tickets

    # ------------------------------------------------------------ group order
    def build_group_orders(self):
        r = self.rng
        baks, proofs = [], []
        bak_id = 0
        for p in PERIODS:
            days = self.days(p)
            for i in range(20):
                bak_id += 1
                d = r.choice(days[:-3])
                st = r.choice(list(STATIONS))
                dest = r.choice([s for s in STATIONS if s != st])
                company = self.fake.company()
                rate = r.choice((10, 15, 20))
                shift = r.choice((1, 2))
                t = self.rand_time(shift)
                op = self.roster[(d, st, shift)][1]["employee_name"]
                order = self.new_order_no(d)
                dep = (d + timedelta(days=r.randint(3, 14)), time(r.randint(6, 20), 0))
                n = r.randint(10, 40)
                tks = [self.make_ticket(d=d, t=t, origin=st, dest=dest, channel="B2B", pay="TRANSFER",
                                        order_no=order, pax=self.passenger(), operator=op, shift=shift, window=1,
                                        station=st, discount_type="GROUP", discount_rate=rate,
                                        seat_class="PREMIUM_ECONOMY", dep=dep, bak_id=bak_id, partner=company)
                       for _ in range(n)]
                self.tickets.extend(tks)
                total = sum(tk["after_tax_price"] for tk in tks)
                approver = self.station_staff[st]["kepala"]
                bak_no = f"BAK/{st}/{d:%Y%m}/{bak_id:03d}"
                baks.append({"bak_id": bak_id, "no_dokumen": bak_no, "periode": p, "tanggal_dokumen": d.isoformat(),
                             "nama_pihak_kedua": company, "order_no": order, "jumlah_tiket": n,
                             "total_nominal": total, "discount_type": "GROUP", "discount_rate": rate,
                             "stasiun": st, "disetujui_oleh": approver["employee_id"], "bukti_transfer": "1",
                             "manifest": "1", "status": "final", "keterangan": None})
                tdate = max(d - timedelta(days=r.randint(0, 2)), date(2025, 1, 1))
                conf = round(r.uniform(0.72, 0.99), 3)
                proofs.append({"proof_id": f"PRF{bak_id:04d}", "bak_no": bak_no,
                               "file_name": f"bukti_{bak_id:04d}.pdf",
                               "file_sha256": sha256_bytes(f"proof-{bak_id}".encode()),
                               "uploaded_at": f"{d.isoformat()} {r.randint(8, 16):02d}:{r.randint(0, 59):02d}:00",
                               "sender_name": company, "sender_bank": r.choice(BANKS),
                               "sender_account_masked": f"****{r.randint(1000, 9999)}",
                               "transfer_date": tdate.isoformat(), "amount": total,
                               "reference_no": f"FT{tdate:%y%m%d}{r.randint(10000, 99999)}",
                               "ocr_confidence": conf,
                               "a2i_status": "REVIEWED" if conf < 0.85 else "AUTO",
                               "a2i_maker": r.choice(self.it_staff)["employee_id"] if conf < 0.85 else None,
                               "a2i_checker": self.station_staff[st]["supervisor"][0]["employee_id"] if conf < 0.85 else None,
                               "periode": p})
        self.baks, self.proofs = baks, proofs

    # ------------------------------------------------------------------ refunds
    def build_refunds(self):
        r = self.rng
        refunds = []
        candidates = [tk for tk in self.tickets if tk["ticketing_channel"] != "B2B"
                      and tk["departure_date"] <= "2025-03-28"]
        r.shuffle(candidates)
        per_station = {}
        for tk in candidates[:240]:
            st = tk["origin_code"]
            dep = date.fromisoformat(tk["departure_date"])
            pur = date.fromisoformat(tk["purchase_date"])
            rdate = pur + timedelta(days=r.randint(0, (dep - pur).days))
            per_station[st] = per_station.get(st, 0) + 1
            amount = half_up(tk["after_tax_price"] * REFUND_RATE)
            cash = tk["nationality"] == "IDN" and amount <= CASH_REFUND_LIMIT and r.random() < 0.3
            refunds.append(self.refund_row(tk, st, rdate, amount, "CASH" if cash else "TRANSFER", per_station[st]))
        self.refunds = refunds

    def refund_row(self, tk: dict, st: str, rdate: date, amount: int, method: str, seq: int) -> dict:
        r = self.rng
        conf = round(r.uniform(0.70, 0.99), 3)
        return {"refund_id": f"RF-{st}-{rdate:%y%m}-{seq:04d}", "station_code": st,
                "refund_date": rdate.isoformat(), "ticket_no": tk["ticket_no"], "order_no": tk["order_no"],
                "passenger_name": tk["passenger_name"], "nationality": tk["nationality"],
                "id_number": tk["nik_passport_no"], "refund_method": method,
                "bank_name": r.choice(BANKS) if method == "TRANSFER" else None,
                "bank_account": f"{r.randint(10**9, 10**10 - 1)}" if method == "TRANSFER" else None,
                "account_holder": tk["passenger_name"] if method == "TRANSFER" else None,
                "ticket_price": tk["after_tax_price"], "refund_amount": amount,
                "processed_by": self.roster[(rdate, st, 1)][1]["employee_name"],
                "approved_by": self.station_staff[st]["kepala"]["employee_name"],
                "ocr_confidence": conf, "a2i_status": "REVIEWED" if conf < 0.85 else "AUTO",
                "a2i_maker": r.choice(self.it_staff)["employee_id"] if conf < 0.85 else None,
                "a2i_checker": self.station_staff[st]["supervisor"][1]["employee_id"] if conf < 0.85 else None}

    # --------------------------------------------------------------- anomalies
    def inject_ticket_anomalies(self):
        r = self.rng
        by_id = {tk["ticket_no"]: tk for tk in self.tickets}

        def pick(pred, k):
            pool = [tk for tk in self.tickets if pred(tk) and tk["ticket_no"] not in self.injected_tickets]
            chosen = r.sample(pool, k)
            self.injected_tickets.update(tk["ticket_no"] for tk in chosen)
            return chosen

        # R05 diskon titipan: diskon GROUP di loket tanpa BAK
        orders = sorted({tk["order_no"] for tk in self.tickets if tk["ticketing_channel"] == "COUNTER"})
        multi = [o for o in orders if sum(1 for tk in self.tickets if tk["order_no"] == o) >= 3]
        for o in r.sample(multi, 2):
            for tk in [tk for tk in self.tickets if tk["order_no"] == o][:3]:
                tk["discount_type"], tk["discount_rate"] = "GROUP", 25
                tk["before_tax_price"] = half_up(tk["original_ticket_price"] * 0.75)
                tk["after_tax_price"] = half_up(tk["before_tax_price"] * (1 + TAX_RATE))
                self.injected_tickets.add(tk["ticket_no"])
                self.add_truth("R05_UNAUTHORIZED_DISCOUNT", "ticket", tk["ticket_no"], "IT", tk["purchase_date"][:7],
                               f"Diskon GROUP 25% di loket tanpa BAK, operator {tk['operator_name']}")

        # R07 operator tidak terjadwal di shift tersebut
        for tk in pick(lambda t: t["ticketing_channel"] == "COUNTER", 5):
            other = r.choice([s for s in STATIONS if s != tk["ticketing_station"]])
            tk["operator_name"] = r.choice(self.station_staff[other]["petugas"])["employee_name"]
            self.add_truth("R07_OPERATOR_NOT_ON_SHIFT", "ticket", tk["ticket_no"], "IT", tk["purchase_date"][:7],
                           "Operator tidak ada di jadwal shift HCM untuk stasiun, tanggal, dan shift tersebut")

        # R08 kesalahan perhitungan harga
        for tk in pick(lambda t: t["ticketing_channel"] in ("APP", "WEB", "TVM"), 4):
            tk["before_tax_price"] += r.choice((5_000, 10_000, 25_000))
            self.add_truth("R08_PRICE_CALC_ERROR", "ticket", tk["ticket_no"], "IT", tk["purchase_date"][:7],
                           "before_tax_price tidak konsisten dengan harga asli, diskon, dan pajak")

        # Group order: R06 diskon tidak sesuai BAK (+R04), R04 tiket melebihi BAK
        bak_r06, bak_r04 = r.sample(self.baks, 2)
        tks = [tk for tk in self.tickets if tk["bak"] == bak_r06["bak_id"]][:3]
        for tk in tks:
            tk["discount_rate"] = 30
            tk["before_tax_price"] = half_up(tk["original_ticket_price"] * 0.70)
            tk["after_tax_price"] = half_up(tk["before_tax_price"] * (1 + TAX_RATE))
            self.injected_tickets.add(tk["ticket_no"])
            self.add_truth("R06_DISCOUNT_NOT_PER_BAK", "ticket", tk["ticket_no"], "IT", bak_r06["periode"],
                           f"Diskon 30% padahal BAK {bak_r06['no_dokumen']} menyetujui {bak_r06['discount_rate']}%")
        self.add_truth("R04_BAK_TICKET_MISMATCH", "bak", bak_r06["no_dokumen"], "FINANCE", bak_r06["periode"],
                       "Total nilai tiket di export berbeda dengan total BAK (akibat diskon tidak sesuai BAK)")
        base = next(tk for tk in self.tickets if tk["bak"] == bak_r04["bak_id"])
        for _ in range(2):
            extra = dict(base)
            self.ticket_seq += 1
            extra["ticket_no"] = f"TK{self.ticket_seq:010d}"
            pax = self.passenger()
            extra.update(passenger_name=pax["name"], nik_passport_no=pax["id_no"], nationality=pax["nationality"])
            self.tickets.append(extra)
            self.injected_tickets.add(extra["ticket_no"])
        self.add_truth("R04_BAK_TICKET_MISMATCH", "bak", bak_r04["no_dokumen"], "FINANCE", bak_r04["periode"],
                       "Export berisi 2 tiket lebih banyak dari jumlah tiket di BAK")
        self.skip_baks = {bak_r06["bak_id"], bak_r04["bak_id"]}

        # Pengecoh R07: tukar shift yang tercatat di HCM. Petugas stasiun lain menggantikan dan jadwalnya
        # diperbarui, jadi transaksinya sah. Satu nama ditulis huruf kecil dengan spasi berlebih.
        for i, tk in enumerate(pick(lambda t: t["ticketing_channel"] == "COUNTER", 2)):
            other = r.choice([s for s in STATIONS if s != tk["ticketing_station"]])
            sub = r.choice(self.station_staff[other]["petugas"])
            self.roster_rows.append({"roster_date": tk["purchase_date"], "station_code": tk["ticketing_station"],
                                     "shift_no": tk["shift_no"],
                                     "shift_start": "05:30" if tk["shift_no"] == 1 else "13:30",
                                     "shift_end": "13:30" if tk["shift_no"] == 1 else "21:30",
                                     "employee_id": sub["employee_id"], "role": "PETUGAS_LOKET_PENGGANTI",
                                     "window_no": tk["window_no"]})
            tk["operator_name"] = sub["employee_name"] if i == 0 else f"  {sub['employee_name'].lower()} "
            self.add_decoy("R07_OPERATOR_NOT_ON_SHIFT", "ticket", tk["ticket_no"],
                           "Petugas pengganti dengan tukar shift yang tercatat di jadwal HCM")

        # Pengecoh R08: selisih pembulatan 1 rupiah, masih dalam toleransi kebijakan.
        for tk in pick(lambda t: t["ticketing_channel"] in ("APP", "WEB", "TVM"), 2):
            tk["before_tax_price"] += 1
            self.add_decoy("R08_PRICE_CALC_ERROR", "ticket", tk["ticket_no"],
                           "Selisih pembulatan Rp1 pada harga sebelum pajak")

    def inject_duplicates(self):
        r = self.rng
        pool = [tk for tk in self.tickets if tk["ticketing_channel"] == "APP" and tk["ticket_no"] not in self.injected_tickets]
        for tk in r.sample(pool, 2):
            self.tickets.append(dict(tk))
            self.injected_tickets.add(tk["ticket_no"])
            self.add_truth("R09_DUPLICATE_TICKET", "ticket", tk["ticket_no"], "IT", tk["purchase_date"][:7],
                           "ticket_no muncul lebih dari sekali di export")

    def inject_refund_anomalies(self):
        r = self.rng
        missing = ("STC", "2025-03")  # laporan refund STC Maret tidak diunggah (lihat U02)
        ok = [x for x in self.refunds if (x["station_code"], x["refund_date"][:7]) != missing]
        used = set()

        def pick(pred, k):
            pool = [x for x in ok if pred(x) and x["refund_id"] not in used]
            chosen = r.sample(pool, k)
            used.update(x["refund_id"] for x in chosen)
            return chosen

        for x in pick(lambda x: x["nationality"] != "IDN" and x["refund_method"] == "TRANSFER", 4):
            x.update(refund_method="CASH", bank_name=None, bank_account=None, account_holder=None)
            self.add_truth("R11_CASH_REFUND_POLICY", "refund", x["refund_id"], "STASIUN", x["refund_date"][:7],
                           f"Refund tunai kepada WNA ({x['nationality']})")
        for x in pick(lambda x: x["refund_method"] == "TRANSFER", 2):
            x["ticket_no"] = f"TK9{r.randint(10**8, 10**9 - 1):09d}"
            self.add_truth("R12_REFUND_TICKET_NOT_FOUND", "refund", x["refund_id"], "STASIUN", x["refund_date"][:7],
                           "Nomor tiket di laporan refund tidak ada di export ticketing")
        for x in pick(lambda x: x["refund_method"] == "TRANSFER", 2):
            x["refund_amount"] = x["ticket_price"]
            self.add_truth("R13_REFUND_AMOUNT_EXCEEDS", "refund", x["refund_id"], "STASIUN", x["refund_date"][:7],
                           "Refund 100% padahal kebijakan maksimal 75% harga tiket")
        self.refund_no_debit = set()
        for x in pick(lambda x: x["refund_method"] == "TRANSFER", 3):
            self.refund_no_debit.add(x["refund_id"])
            self.add_truth("R14_REFUND_NO_BANK_DEBIT", "refund", x["refund_id"], "STASIUN", x["refund_date"][:7],
                           "Refund transfer tidak ditemukan di mutasi debit rekening koran")
        self.missing_refund = missing

        # Pengecoh R11: refund tunai WNI tepat di batas Rp300.000.
        for x in pick(lambda x: x["nationality"] == "IDN" and x["refund_method"] == "TRANSFER"
                      and half_up(x["ticket_price"] * REFUND_RATE) >= CASH_REFUND_LIMIT, 2):
            x.update(refund_method="CASH", refund_amount=CASH_REFUND_LIMIT,
                     bank_name=None, bank_account=None, account_holder=None)
            self.add_decoy("R11_CASH_REFUND_POLICY", "refund", x["refund_id"],
                           "Refund tunai WNI tepat Rp300.000, sesuai kebijakan")
        # Pengecoh R14: nama pemilik rekening memakai tanda hubung, bank menulis tanpa tanda baca,
        # dan debit baru tercatat di hari terakhir jendela pencocokan.
        self.refund_name_variant = set()
        for x in pick(lambda x: x["refund_method"] == "TRANSFER" and len(x["account_holder"].split()) >= 2, 2):
            x["account_holder"] = x["account_holder"].replace(" ", "-", 1)
            self.refund_name_variant.add(x["refund_id"])
            self.add_decoy("R14_REFUND_NO_BANK_DEBIT", "refund", x["refund_id"],
                           "Nama di laporan refund memakai tanda hubung, bank menulis tanpa tanda baca, debit H+3")

    # ------------------------------------------------------------------- bank
    def build_bank(self):
        r = self.rng
        tx = []
        last_day = date(2025, 3, 31)
        fake_proofs = set(x["proof_id"] for x in r.sample([p for p in self.proofs if p["proof_id"] not in
                                                           {f"PRF{b:04d}" for b in self.skip_baks}], 5))
        fake_ids, short_ids = sorted(fake_proofs)[:3], sorted(fake_proofs)[3:]
        # Pengecoh R01/R03: bank tidak mencantumkan nomor referensi pengirim dan kredit baru masuk H+3.
        skip_proofs = fake_proofs | {f"PRF{b:04d}" for b in self.skip_baks}
        noref_ids = set(x["proof_id"] for x in r.sample([p for p in self.proofs if p["proof_id"] not in skip_proofs], 2))
        for p in self.proofs:
            if p["proof_id"] in noref_ids:
                d = min(date.fromisoformat(p["transfer_date"]) + timedelta(days=3), last_day)
                tx.append({"date": d, "desc": f"TRF DR {p['sender_name'].upper()}", "debit": 0,
                           "credit": p["amount"], "truth": None, "decoy": ("R01_FAKE_TRANSFER_PROOF", p["proof_id"])})
                continue
            if p["proof_id"] in fake_ids:
                self.add_truth("R01_FAKE_TRANSFER_PROOF", "proof", p["proof_id"], "FINANCE", p["periode"],
                               "Bukti transfer tidak punya pasangan kredit di rekening koran")
                continue
            amount = p["amount"]
            if p["proof_id"] in short_ids:
                amount -= r.choice((500_000, 1_000_000, 2_500_000))
                self.add_truth("R02_PROOF_AMOUNT_MISMATCH", "proof", p["proof_id"], "FINANCE", p["periode"],
                               "Nominal di bukti transfer lebih besar dari kredit yang benar-benar masuk")
            tx.append({"date": date.fromisoformat(p["transfer_date"]), "desc": f"TRF {p['reference_no']} {p['sender_name'].upper()}",
                       "debit": 0, "credit": amount, "truth": None})
        # setoran tunai harian per stasiun
        cash = {}
        settle = {}
        for tk in self.tickets:
            d = date.fromisoformat(tk["purchase_date"])
            if tk["payment_method"] == "CASH":
                cash[(tk["ticketing_station"], d)] = cash.get((tk["ticketing_station"], d), 0) + tk["after_tax_price"]
            elif tk["payment_gateway"]:
                settle[(tk["payment_gateway"], d)] = settle.get((tk["payment_gateway"], d), 0) + tk["after_tax_price"]
        # Pengecoh R03: setoran tunai penjualan suatu hari baru disetor keesokan harinya.
        late_cash = set(r.sample([k for k in sorted(cash) if k[1] < last_day], 2))
        for (st, d), amt in sorted(cash.items()):
            if (st, d) in late_cash:
                tx.append({"date": d + timedelta(days=1), "desc": f"SETORAN TUNAI {st} {d:%d/%m/%y}", "debit": 0,
                           "credit": amt, "truth": None, "decoy": ("R03_UNMATCHED_BANK_CREDIT", None)})
                continue
            tx.append({"date": d, "desc": f"SETORAN TUNAI {st} {d:%d/%m/%y}", "debit": 0, "credit": amt, "truth": None})
        for (gw, d), amt in sorted(settle.items()):
            if d + timedelta(days=1) <= last_day:
                tx.append({"date": d + timedelta(days=1), "desc": f"STLM {gw} {d:%d%m%y}", "debit": 0, "credit": amt, "truth": None})
        # bunga dan biaya bulanan
        for p in PERIODS:
            end = self.days(p)[-1]
            tx.append({"date": end, "desc": "BUNGA JASA GIRO", "debit": 0, "credit": r.randint(2, 9) * 1_000_000, "truth": None})
            tx.append({"date": end, "desc": "BIAYA ADM REKENING", "debit": 25_000, "credit": 0, "truth": None})
        # debit refund transfer
        for x in self.refunds:
            if x["refund_method"] != "TRANSFER" or x["refund_id"] in self.refund_no_debit:
                continue
            if x["refund_id"] in self.refund_name_variant:
                d = min(date.fromisoformat(x["refund_date"]) + timedelta(days=3), last_day)
                holder = x["account_holder"].replace("-", " ")
            else:
                d = min(date.fromisoformat(x["refund_date"]) + timedelta(days=r.randint(0, 2)), last_day)
                holder = x["account_holder"]
            tx.append({"date": d, "desc": f"TRF KE {holder.upper()} {x['bank_name'].upper()}",
                       "debit": x["refund_amount"], "credit": 0, "truth": None})
        # R03 kredit tanpa sumber jelas
        for p in ["2025-01", "2025-02", "2025-02", "2025-03"]:
            d = r.choice(self.days(p))
            tx.append({"date": d, "desc": f"TRF FT{d:%y%m%d}{r.randint(10000, 99999)} {self.unique_name().upper()}",
                       "debit": 0, "credit": r.randint(20, 150) * 100_000, "truth": "R03"})
        tx.sort(key=lambda x: (x["date"], x["desc"]))
        balance = 5_000_000_000
        rows = []
        seq = {}
        for x in tx:
            seq[x["date"]] = seq.get(x["date"], 0) + 1
            ref = f"BRF{x['date']:%y%m%d}{seq[x['date']]:04d}"
            balance += x["credit"] - x["debit"]
            rows.append({"Tanggal": x["date"].strftime("%d/%m/%Y"), "Tgl Valuta": x["date"].strftime("%d/%m/%Y"),
                         "No Ref Bank": ref, "Keterangan": x["desc"], "Cabang": "0101",
                         "Debet": rupiah_id(x["debit"]), "Kredit": rupiah_id(x["credit"]),
                         "Saldo": rupiah_id(balance), "_period": f"{x['date']:%Y-%m}"})
            if x.get("decoy"):
                rule, eid = x["decoy"]
                if eid:
                    self.add_decoy("R03_UNMATCHED_BANK_CREDIT", "bank_mutation", ref,
                                   "Kredit sah tanpa nomor referensi pengirim, masuk H+3")
                    self.add_decoy(rule, "proof", eid, "Bukti transfer sah, bank tidak mencantumkan referensinya")
                else:
                    self.add_decoy(rule, "bank_mutation", ref, "Setoran tunai penjualan kemarin yang baru disetor hari ini")
            if x["truth"] == "R03":
                self.add_truth("R03_UNMATCHED_BANK_CREDIT", "bank_mutation", ref, "FINANCE", f"{x['date']:%Y-%m}",
                               "Kredit masuk tanpa pasangan bukti transfer, setoran tunai, atau settlement")
        self.bank_rows = rows

    # ----------------------------------------------------------------- recaps
    def build_recaps(self):
        """Rekap stasiun diturunkan dari export (scope: loket dan B2B di stasiun itu), lalu disisipi selisih."""
        r = self.rng
        df = pd.DataFrame(self.tickets).drop_duplicates("ticket_no")
        df = df[df["ticketing_channel"].isin(["COUNTER", "B2B"])]
        self.recaps = {}
        for st in STATIONS:
            sdf = df[df["ticketing_station"] == st]
            orders = (sdf.groupby("order_no")
                      .agg(purchase_date=("purchase_date", "first"), shift_no=("shift_no", "first"),
                           window_no=("window_no", "first"), buyer=("passenger_name", "first"),
                           partner=("b2b_partner", "first"), pax=("ticket_no", "count"),
                           amount=("after_tax_price", "sum"), pay=("payment_method", "first"))
                      .reset_index().sort_values(["purchase_date", "order_no"]))
            self.recaps[st] = orders
        # Stasiun A: 2 nominal salah + 1 order hilang
        a = self.recaps["STA"]
        idx = r.sample(list(a.index), 3)
        for i in idx[:2]:
            a.loc[i, "amount"] -= r.choice((100_000, 250_000))
            self.add_truth("R10_RECAP_MISMATCH", "recap", f"STA|{a.loc[i, 'order_no']}", "STASIUN",
                           a.loc[i, "purchase_date"][:7], "Nominal order di rekap stasiun berbeda dengan export")
        self.add_truth("R10_RECAP_MISMATCH", "recap", f"STA|{a.loc[idx[2], 'order_no']}", "STASIUN",
                       a.loc[idx[2], "purchase_date"][:7], "Order ada di export tetapi tidak ada di rekap stasiun")
        self.recaps["STA"] = a.drop(idx[2])
        # Stasiun C: 1 order hilang + 1 order fiktif
        c = self.recaps["STC"]
        i = r.choice(list(c.index))
        self.add_truth("R10_RECAP_MISMATCH", "recap", f"STC|{c.loc[i, 'order_no']}", "STASIUN",
                       c.loc[i, "purchase_date"][:7], "Order ada di export tetapi tidak ada di rekap stasiun")
        c = c.drop(i)
        ghost = c.iloc[r.randrange(len(c))].copy()
        ghost["order_no"] = f"ORD{ghost['purchase_date'].replace('-', '')}9{r.randint(1000, 9999)}"
        c = pd.concat([c, ghost.to_frame().T]).sort_values(["purchase_date", "order_no"])
        self.add_truth("R10_RECAP_MISMATCH", "recap", f"STC|{ghost['order_no']}", "STASIUN",
                       ghost["purchase_date"][:7], "Order di rekap stasiun tidak ada di export")
        self.recaps["STC"] = c
        # Stasiun B: rekap per shift, 2 total shift salah
        b = self.recaps["STB"].copy()
        b["cash"] = b["amount"].where(b["pay"] == "CASH", 0)
        b["noncash"] = b["amount"] - b["cash"]
        shifts = (b.groupby(["purchase_date", "shift_no"])
                  .agg(pax=("pax", "sum"), cash=("cash", "sum"), noncash=("noncash", "sum"), amount=("amount", "sum"))
                  .reset_index())
        for i in r.sample(list(shifts.index), 2):
            delta = r.choice((150_000, 300_000))
            shifts.loc[i, "cash"] -= delta
            shifts.loc[i, "amount"] -= delta
            self.add_truth("R10_RECAP_MISMATCH", "recap",
                           f"STB|{shifts.loc[i, 'purchase_date']}|{int(shifts.loc[i, 'shift_no'])}", "STASIUN",
                           shifts.loc[i, "purchase_date"][:7], "Total shift di rekap stasiun berbeda dengan export")
        self.recaps["STB"] = shifts

    # ------------------------------------------------------------------ write
    def write_all(self):
        out = self.out
        for sub in ("it", "stasiun", "finance", "hcm", "portal"):
            (out / sub).mkdir(parents=True, exist_ok=True)
        self.files: dict[tuple, Path] = {}

        tdf = pd.DataFrame(self.tickets)
        for p in PERIODS:
            part = tdf[tdf["purchase_date"].str.startswith(p)].sort_values(["purchase_date", "purchase_time", "ticket_no"]).copy()
            part["seq_no"] = range(1, len(part) + 1)
            path = out / "it" / f"export_ticketing_{p}.csv"
            part.to_csv(path, index=False)
            self.files[("IT", "EXPORT_TICKETING", "-", p)] = path

        for p in PERIODS:
            y, m = map(int, p.split("-"))
            # Stasiun A: xlsx satu baris per order, tanggal teks dd/mm/yyyy, baris TOTAL di bawah
            a = self.recaps["STA"][self.recaps["STA"]["purchase_date"].str.startswith(p)]
            wb = Workbook()
            ws = wb.active
            ws.title = "Rekap"
            ws.append(["No", "Tanggal", "Shift", "No. Order", "Nama Pemesan", "Jml Tiket", "Metode Bayar", "Total (Rp)"])
            for n, row in enumerate(a.itertuples(), 1):
                ws.append([n, date.fromisoformat(row.purchase_date).strftime("%d/%m/%Y"), int(row.shift_no), row.order_no,
                           row.partner if isinstance(row.partner, str) else row.buyer, int(row.pax), row.pay, int(row.amount)])
            ws.append([None, None, None, None, "TOTAL", int(a["pax"].sum()), None, int(a["amount"].sum())])
            path = out / "stasiun" / f"rekap_STA_{p}.xlsx"
            wb.save(path)
            self.files[("STASIUN", "REKAP_PENJUALAN", "STA", p)] = path
            # Stasiun B: csv titik koma per shift, nominal teks "Rp1.234.000"
            b = self.recaps["STB"][self.recaps["STB"]["purchase_date"].str.startswith(p)]
            fmt = lambda v: "Rp" + f"{int(v):,}".replace(",", ".")
            path = out / "stasiun" / f"rekap_STB_{p}.csv"
            with path.open("w", newline="", encoding="utf-8") as fh:
                w = csv.writer(fh, delimiter=";")
                w.writerow(["Tgl", "Shift", "Jml Tiket", "Tunai", "Non Tunai", "Total"])
                for row in b.itertuples():
                    w.writerow([date.fromisoformat(row.purchase_date).strftime("%d-%m-%Y"),
                                "Pagi" if int(row.shift_no) == 1 else "Siang", int(row.pax),
                                fmt(row.cash), fmt(row.noncash), fmt(row.amount)])
            self.files[("STASIUN", "REKAP_PENJUALAN", "STB", p)] = path
            # Stasiun C: xlsx dengan baris judul, header bahasa Inggris, tanggal sel Excel
            c = self.recaps["STC"][self.recaps["STC"]["purchase_date"].str.startswith(p)]
            wb = Workbook()
            ws = wb.active
            ws.title = f"{MONTH_ID[m]}{y}"
            ws.append(["REKAP PENJUALAN TIKET STASIUN C"])
            ws.append([f"Periode: {MONTH_NAME[m]} {y}"])
            ws.append([])
            ws.append(["Order Number", "Sales Date", "Window", "Pax", "Amount", "Payment", "Remarks"])
            for row in c.itertuples():
                ws.append([row.order_no, date.fromisoformat(row.purchase_date),
                           int(row.window_no) if pd.notna(row.window_no) else None, int(row.pax), int(row.amount),
                           row.pay, "group" if isinstance(row.partner, str) else None])
                ws.cell(ws.max_row, 2).number_format = "DD-MMM-YY"
            path = out / "stasiun" / f"rekap_STC_{p}.xlsx"
            wb.save(path)
            self.files[("STASIUN", "REKAP_PENJUALAN", "STC", p)] = path

        rdf = pd.DataFrame(self.refunds)
        for p in PERIODS:
            for st in STATIONS:
                if (st, p) == self.missing_refund:
                    continue
                part = rdf[(rdf["station_code"] == st) & rdf["refund_date"].str.startswith(p)]
                path = out / "stasiun" / f"refund_{st}_{p}.csv"
                part.to_csv(path, index=False)
                self.files[("STASIUN", "LAPORAN_REFUND", st, p)] = path

        bdf = pd.DataFrame(self.baks)
        pdf = pd.DataFrame(self.proofs)
        kdf = pd.DataFrame(self.bank_rows)
        for p in PERIODS:
            path = out / "finance" / f"bak_group_order_{p}.csv"
            bdf[bdf["periode"] == p].to_csv(path, index=False)
            self.files[("FINANCE", "BAK_GROUP_ORDER", "-", p)] = path
            path = out / "finance" / f"bukti_transfer_{p}.csv"
            pdf[pdf["periode"] == p].drop(columns="periode").to_csv(path, index=False)
            self.files[("FINANCE", "BUKTI_TRANSFER", "-", p)] = path
            path = out / "finance" / f"rekening_koran_{p}.csv"
            kdf[kdf["_period"] == p].drop(columns="_period").to_csv(path, index=False)
            self.files[("FINANCE", "REKENING_KORAN", "-", p)] = path

        pd.DataFrame(self.employees).to_csv(out / "hcm" / "employees.csv", index=False)
        pd.DataFrame(self.roster_rows).to_csv(out / "hcm" / "shift_roster.csv", index=False)

    # ------------------------------------------------------------- upload log
    def build_upload_log(self):
        r = self.rng
        obligations = [("STASIUN", "REKAP_PENJUALAN", st) for st in STATIONS] + \
                      [("STASIUN", "LAPORAN_REFUND", st) for st in STATIONS] + \
                      [("FINANCE", "REKENING_KORAN", "-"), ("FINANCE", "BUKTI_TRANSFER", "-"),
                       ("FINANCE", "BAK_GROUP_ORDER", "-"), ("IT", "EXPORT_TICKETING", "-")]
        keys = [(dept, src, st, p) for p in PERIODS for dept, src, st in obligations]
        missing = ("STASIUN", "LAPORAN_REFUND", "STC", "2025-03")
        rest = [k for k in keys if k != missing]
        chosen = r.sample(rest, 8)
        late, failed, pending, excessive, tampered = chosen[:3], chosen[3:4], chosen[4:6], chosen[6:7], chosen[7:8]
        edge = r.choice([k for k in rest if k not in chosen])  # pengecoh U01: lolos validasi tepat di detik tenggat
        uid = 0
        rows = []
        for k in keys:
            dept, src, st, p = k
            sid = "|".join(k)
            if k == missing:
                self.add_truth("U02_MISSING_UPLOAD", "submission", sid, dept, p, "Kewajiban unggah periode ini tidak dipenuhi")
                continue
            y, m = map(int, p.split("-"))
            deadline = datetime(y + (m == 12), m % 12 + 1, 5, 23, 59, 59)
            if dept == "STASIUN":
                uploader = self.station_staff[st]["supervisor"][2]["employee_id"]
                owner = self.station_staff[st]["kepala"]["employee_id"]
            elif dept == "FINANCE":
                uploader, owner = self.finance_staff[0]["employee_id"], self.finance_mgr["employee_id"]
            else:
                uploader, owner = self.it_staff[0]["employee_id"], self.it_mgr["employee_id"]
            path = self.files[k]
            nrows = max(sum(1 for _ in path.open(encoding="utf-8", errors="ignore")) - 1, 0) if path.suffix == ".csv" else None
            if path.suffix == ".xlsx":
                nrows = len(pd.read_excel(path, header=None))
            if k in excessive:
                statuses = ["FAILED", "PASSED", "PASSED", "PASSED"]
            elif k in failed:
                statuses = ["FAILED", "FAILED"]
            elif r.random() < 0.25:
                statuses = ["FAILED", "PASSED"]
            else:
                statuses = ["PASSED"]
            t = datetime(y + (m == 12), m % 12 + 1, 1, 8) + timedelta(minutes=r.randint(0, 3 * 24 * 60))
            if k in late:
                t = deadline + timedelta(hours=r.randint(6, 72))
            if k == edge:
                statuses, t = ["PASSED"], deadline
                self.add_decoy("U01_LATE_UPLOAD", "submission", sid, "Lolos validasi tepat pada detik tenggat")
            for rev, status in enumerate(statuses, 1):
                uid += 1
                final = rev == len(statuses)
                checksum = sha256_file(path) if final else sha256_bytes(f"{sid}-{rev}".encode())
                if final and k in tampered:
                    checksum = sha256_bytes(f"{sid}-original".encode())
                approval, approved_by, approved_at = None, None, None
                if status == "PASSED" and not final:
                    approval = "SUPERSEDED"
                elif status == "PASSED" and final:
                    if k in pending:
                        approval = "PENDING"
                    else:
                        approval, approved_by = "APPROVED", owner
                        approved_at = (t + timedelta(hours=r.randint(1, 20))).isoformat(sep=" ")
                rows.append({"upload_id": f"UP{uid:05d}", "department": dept, "source_type": src,
                             "station_code": st, "period": p, "file_name": path.name,
                             "object_key": f"landing/{dept.lower()}/{src.lower()}/{p}/{path.name}",
                             "uploader_id": uploader, "uploaded_at": t.isoformat(sep=" "),
                             "deadline_at": deadline.isoformat(sep=" "), "revision_no": rev,
                             "validation_status": status,
                             "validation_errors": "kolom wajib kosong; format tanggal tidak sesuai template" if status == "FAILED" else None,
                             "row_count": nrows, "checksum_sha256": checksum, "approval_status": approval,
                             "approved_by": approved_by, "approved_at": approved_at})
                t += timedelta(hours=r.randint(2, 10))
            if k in late:
                self.add_truth("U01_LATE_UPLOAD", "submission", sid, dept, p, "Unggahan pertama yang lolos validasi melewati tenggat")
            if k in failed:
                self.add_truth("U03_VALIDATION_FAILED", "submission", sid, dept, p, "Revisi terakhir masih gagal validasi template")
            if k in pending:
                self.add_truth("U04_NOT_APPROVED", "submission", sid, dept, p, "Unggahan lolos validasi tetapi belum disetujui data owner")
            if k in excessive:
                self.add_truth("U05_EXCESSIVE_REVISION", "submission", sid, dept, p, "Jumlah revisi melebihi batas 3")
            if k in tampered:
                self.add_truth("U06_CHECKSUM_MISMATCH", "submission", sid, dept, p, "Checksum objek di storage berbeda dengan checksum saat unggah")
        pd.DataFrame(rows).to_csv(self.out / "portal" / "upload_log.csv", index=False)

    def run(self):
        self.build_hcm()
        self.build_tickets()
        self.build_group_orders()
        self.inject_ticket_anomalies()
        self.build_refunds()
        self.inject_refund_anomalies()
        self.build_bank()
        self.build_recaps()
        self.inject_duplicates()
        self.write_all()
        self.build_upload_log()
        truth = pd.DataFrame(self.truth)
        truth.to_csv(self.out / "anomalies_truth.csv", index=False)
        pd.DataFrame(self.decoys).to_csv(self.out / "decoys_truth.csv", index=False)
        print(f"Tiket: {len(self.tickets)}, BAK: {len(self.baks)}, refund: {len(self.refunds)}, "
              f"mutasi bank: {len(self.bank_rows)}, pegawai: {len(self.employees)}")
        print(f"Anomali disisipkan: {len(truth)}")
        print(truth["rule_code"].value_counts().sort_index().to_string())
        print(f"Kasus pengecoh (harus lolos tanpa ditandai): {len(self.decoys)}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=str(Path(__file__).parent / "output"))
    ap.add_argument("--seed", type=int, default=SEED)
    args = ap.parse_args()
    Generator(Path(args.out), args.seed).run()


if __name__ == "__main__":
    main()
